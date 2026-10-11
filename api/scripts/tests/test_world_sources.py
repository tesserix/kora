import json
import subprocess
import sys
from pathlib import Path

import openpyxl


SCRIPTS = Path(__file__).resolve().parents[1]


def test_cofid_joins_nutrients_by_food_code_not_row_order(tmp_path):
    w = openpyxl.Workbook()
    w.remove(w.active)
    sheets = {
        "1.3 Proximates": (
            [
                "Food Code",
                "Food Name",
                "Energy (kcal)",
                "Protein (g)",
                "Carbohydrate (g)",
                "Fat (g)",
                "AOAC fibre (g)",
                "Total sugars (g)",
                "Satd FA /100g fd (g)",
            ],
            [None, None, "KCALS", "PROT", "CHO", "FAT", "AOACFIB", "TOTSUG", "SATFOD"],
            [["A", "UK meal", 100, 5, 15, 2, 1, 0, 1]],
        ),
        "1.4 Inorganics": (
            ["Food Code", "Food Name", "Sodium (mg)"],
            [None, None, "NA"],
            [["B", "Other food", 999], ["A", "UK meal", 420]],
        ),
        "1.5 Vitamins": (
            ["Food Code", "Food Name", "Retinol Equivalent (µg)", "Vitamin C (mg)"],
            [None, None, "RETEQU", "VITC"],
            [["A", "UK meal", 12, "N"]],
        ),
    }
    for name, (header, codes, rows) in sheets.items():
        s = w.create_sheet(name)
        for row in [header, codes, [], *rows]:
            s.append(row)
    path = tmp_path / "cofid.xlsx"
    w.save(path)
    result = subprocess.run(
        [sys.executable, str(SCRIPTS / "cofid_convert.py"), str(path)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr
    food = json.loads(result.stdout)[0]
    assert food["locale"] == "GB"
    assert food["sodium_mg_per_100g"] == 420
    assert food["sugar_per_100g"] == 0
    assert food["vitamin_a_re_ug_per_100g"] == 12
    assert "vitamin_c_mg_per_100g" not in food


def test_mext_keeps_rae_and_estimated_values_without_treating_trace_as_zero(tmp_path):
    w = openpyxl.Workbook()
    s = w.active
    s.title = "Table"
    for _ in range(5):
        s.append([])
    s.append(
        [
            "Food and Description",
            "Energy (kcal)",
            "Protein, calculated from reference nitrogen",
            "Lipid",
            "Carbohydrate, total, calculated by difference",
            "Dietary fiber, total",
            "Sodium",
            "Retinol activity equivalents",
            "Vitamin D",
        ]
    )
    s.append([])
    s.append(
        [
            None,
            "kcal/100 g",
            "g/100 g",
            "g/100 g",
            "g/100 g",
            "g/100 g",
            "mg/100 g",
            "µg/100 g",
            "µg/100 g",
        ]
    )
    s.append(["Japanese meal", 100, "(5)", 2, 15, 1, 420, "(12)", "Tr"])
    path = tmp_path / "japan.xlsx"
    w.save(path)
    result = subprocess.run(
        [sys.executable, str(SCRIPTS / "mext_convert.py"), str(path)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr
    food = json.loads(result.stdout)[0]
    assert food["locale"] == "JP"
    assert food["protein_per_100g"] == 5
    assert food["sodium_mg_per_100g"] == 420
    assert food["vitamin_a_rae_ug_per_100g"] == 12
    assert "vitamin_d_ug_per_100g" not in food
